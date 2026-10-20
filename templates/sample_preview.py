import sys, io
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont
prs = Presentation(sys.argv[1]); out = sys.argv[2]
DPI = 60; E = 914400
W, H = int(prs.slide_width / E * DPI), int(prs.slide_height / E * DPI)
def font(pt, bold=False):
    return ImageFont.truetype("/System/Library/Fonts/AppleSDGothicNeo.ttc", max(6, int(pt / 72 * DPI)), index=6 if bold else 0)
def px(v): return int(v / E * DPI)
def wrap(d, txt, f, w):
    lines, cur = [], ""
    for ch in txt:
        if d.textlength(cur + ch, font=f) > w: lines.append(cur); cur = ch
        else: cur += ch
    return lines + [cur]
for n, sl in enumerate(prs.slides, 1):
    im = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(im)
    for sh in sl.shapes:
        x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        if sh.shape_type == 13:
            im.paste(Image.open(io.BytesIO(sh.image.blob)).convert("RGB").resize((max(w,1), max(h,1))), (x, y)); continue
        if sh.has_table if hasattr(sh, "has_table") else False:
            t = sh.table; cy = y
            for r in t.rows:
                cx = x
                for c, col in enumerate(t.columns):
                    cw = px(col.width); rh = px(r.height)
                    d.rectangle([cx, cy, cx+cw, cy+rh], outline="gray")
                    d.text((cx+4, cy+4), r.cells[c].text, fill="black", font=font(18))
                    cx += cw
                cy += px(r.height)
            continue
        try:
            if sh.fill.type == 1: d.rectangle([x, y, x+w, y+h], fill=tuple(sh.fill.fore_color.rgb))
        except Exception: pass
        if sh.has_text_frame:
            ty = y + 4
            for p in sh.text_frame.paragraphs:
                if not p.runs: continue
                r = p.runs[0]; f = font(r.font.size.pt if r.font.size else 18, r.font.bold)
                col = tuple(r.font.color.rgb) if r.font.color and r.font.color.type else (0,0,0)
                for ln in wrap(d, p.text, f, w - 8):
                    d.text((x+4, ty), ln, fill=col, font=f); ty += int(f.size * 1.25)
                ty += int(f.size * 0.4)
            if ty > y + h + 2 and sh.height < E * 6:
                d.rectangle([x, y, x+w, ty], outline="red", width=2)
    im.save(f"{out}/s{n:02d}.png")

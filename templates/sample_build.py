import re
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

FIG, OUT = sys.argv[1], sys.argv[2]
FONT = "Apple SD Gothic Neo"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ORANGE = RGBColor(0xE0, 0x7A, 0x2E)
LIGHT = RGBColor(0xEE, 0xF3, 0xF8)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x88, 0x88, 0x88)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SRC = "출처: 동아 확률과 통계 Ⅲ-1. 확률분포 (pp. 81–85)"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=22, color=DARK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(10)
        runs(p, line, size, color, bold)
    return tb


def runs(p, line, size, color, bold):
    """'p_{i}' 표기를 아래첨자로 변환해 문단에 추가"""
    for k, part in enumerate(re.split(r"_\{([^}]*)\}", line)):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
        r.font.color.rgb = color
        if k % 2:
            r.font._rPr.set("baseline", "-25000")


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def pic(slide, name, x, y, w=None, h=None):
    """주어진 상자(w×h) 안에 비율을 유지하며 가운데 배치"""
    path = f"{FIG}/{name}.png"
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


page = [0]
# python-pptx가 만드는 발표자 노트는 Keynote에서 열리지 않음 → 노트는 json으로 따로 저장해 Keynote에서 삽입
NOTES = []


def new(title, msg=None, notes=None, source=True):
    s = prs.slides.add_slide(BLANK)
    page[0] += 1
    rect(s, 0.6, 0.55, 0.12, 0.6, ORANGE)
    text(s, 0.85, 0.42, 11.5, 0.8, title, size=30, color=NAVY, bold=True)
    if msg:
        text(s, 0.85, 1.2, 11.8, 0.6, msg, size=20, color=GRAY)
    if source:
        text(s, 0.6, 6.95, 8, 0.4, SRC, size=11, color=GRAY)
    text(s, 12.0, 6.95, 0.8, 0.4, str(page[0]), size=12, color=GRAY, align=PP_ALIGN.RIGHT)
    NOTES.append(notes or "")
    return s


def bullets(slide, x, y, w, items, size=22):
    tb = text(slide, x, y, w, 0.5 * len(items) + 0.5, ["• " + i for i in items], size=size)
    for p in tb.text_frame.paragraphs:
        p.space_after = Pt(16)
    return tb


def keybox(slide, x, y, w, h, head, lines, size=20):
    rect(slide, x, y, w, h, LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.06
    text(slide, x + 0.3, y + 0.2, w - 0.6, 0.5, head, size=size, color=ORANGE, bold=True)
    text(slide, x + 0.3, y + 0.75, w - 0.6, h - 0.9, lines, size=size)


# 1. 표지
s = prs.slides.add_slide(BLANK)
page[0] += 1
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0.9, 2.55, 0.15, 1.9, ORANGE)
text(s, 1.3, 1.6, 11, 0.6, "확률과 통계  |  Ⅲ. 통계 – 1. 확률분포", size=20, color=RGBColor(0xC8, 0xD6, 0xE5))
text(s, 1.3, 2.45, 11, 1.2, "확률변수와 확률분포", size=54, color=WHITE, bold=True)
text(s, 1.3, 3.65, 11, 0.8, "우연한 결과를 수로 바꾸는 방법", size=26, color=ORANGE)
text(s, 1.3, 6.2, 11, 0.5, "공과대학 1학년  ·  01회차", size=18, color=RGBColor(0xC8, 0xD6, 0xE5))
NOTES.append("01회차 시작. 확률분포 단원 전체의 출발점이 되는 개념 소개.")

# 2. 학습목표
s = new("학습목표", notes="오늘의 목표 3가지 안내. 이산과 연속의 구분이 이후 이항분포·정규분포로 이어지는 핵심.")
for i, (head, body) in enumerate([
    ("01", "확률변수의 뜻"),
    ("02", "이산확률변수와 확률질량함수의 성질"),
    ("03", "연속확률변수와 확률밀도함수의 성질"),
]):
    y = 2.0 + i * 1.45
    rect(s, 1.2, y, 1.1, 1.1, ORANGE, MSO_SHAPE.OVAL)
    text(s, 1.2, y, 1.1, 1.1, head, size=28, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.7, y, 9.5, 1.1, body, size=28, color=DARK, anchor=MSO_ANCHOR.MIDDLE)

# 3. 도입: 동전 두 번
s = new("동전 두 번 던지기: 결과를 수로 바꾸기",
        "앞면 H, 뒷면 T  |  X = 앞면이 나온 횟수",
        notes="교과서 81쪽 탐구하기. 표본공간 S = {(H,H), (H,T), (T,H), (T,T)}. 근원사건 4개가 X의 값 2, 1, 0에 대응. 서로 다른 근원사건 (H,T), (T,H)가 같은 값 1에 대응하는 점 강조.")
pic(s, "01_mapping", 0.8, 1.9, 6.6, 4.9)
keybox(s, 7.8, 2.4, 4.9, 3.4, "관찰 포인트", [
    "근원사건마다 실수 1개가 대응",
    "(H, T), (T, H) → 같은 값 1",
    "결과가 '수'가 되는 순간 계산 가능",
])

# 4. 확률변수의 정의
s = new("확률변수 (random variable)",
        "표본공간의 각 근원사건에 하나의 실수를 대응시키는 관계",
        notes="정의는 교과서 81쪽 표기 그대로. 공학 예시 3가지는 자체 예시 (교과서 외). 불량품 개수는 셀 수 있는 값, 수명과 오차는 구간 안의 모든 실수라는 차이로 다음 슬라이드 연결.")
text(s, 0.85, 1.85, 11.5, 0.5, "공학 속 확률변수", size=22, color=NAVY, bold=True)
pic(s, "02_engineering", 0.6, 2.4, 12.1, 4.4)

# 5. 이산 vs 연속
s = new("확률변수의 두 종류: 이산과 연속",
        "핵심 기준은 '취하는 값을 셀 수 있는지 여부'",
        notes="이산: 값이 유한개이거나 자연수처럼 셀 수 있는 경우 (교과서 82쪽). 연속: 어떤 범위에 속하는 모든 실수 값 (교과서 84쪽). 예시로 불량품 개수 vs 대기 시간 대비.")
pic(s, "03_disc_vs_cont", 0.8, 1.9, 7.4, 3.6)
keybox(s, 8.6, 1.95, 4.1, 2.0, "이산확률변수", ["불량품 개수, 고객 수,", "주사위 눈"], size=19)
keybox(s, 8.6, 4.15, 4.1, 2.0, "연속확률변수", ["대기 시간, 부품 수명,", "측정 오차"], size=19)

# 6. 이산확률분포
s = new("이산확률변수의 확률분포",
        "X의 값과 각 값의 확률 사이의 관계 = 확률분포",
        notes="교과서 82쪽 보기. 동전 두 번에서 X의 확률분포. 표와 그래프는 같은 정보의 두 표현. P(X=x)를 확률질량함수라고 부름.")
tbl = s.shapes.add_table(2, 5, Inches(0.9), Inches(2.3), Inches(5.6), Inches(1.4)).table
for c, (a, b) in enumerate(zip(["X", "0", "1", "2", "합계"], ["P(X=x)", "1/4", "1/2", "1/4", "1"])):
    for r, v in enumerate([a, b]):
        cell = tbl.cell(r, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if c == 0 else (LIGHT if r == 0 else WHITE)
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        runs(para, v, 20, WHITE if c == 0 else DARK, c == 0)
keybox(s, 0.9, 4.2, 5.6, 2.3, "확률질량함수", ["P(X = x_{i}) = p_{i}  (i = 1, 2, …, n)", "확률분포를 나타내는 함수"], size=19)
pic(s, "04_coin_pmf", 6.9, 1.9, 5.9, 4.8)

# 7. 성질
s = new("확률질량함수의 성질", "모든 확률은 0과 1 사이, 전체 합은 1",
        notes="교과서 82쪽 성질 1, 2. 띠그래프로 '전체 합 = 1' 시각화. 분포표의 빈칸 문제는 성질 2로 해결.")
keybox(s, 0.9, 2.0, 5.6, 1.6, "성질 1", ["0 ≤ p_{i} ≤ 1"], size=24)
keybox(s, 6.9, 2.0, 5.6, 1.6, "성질 2", ["p_{1} + p_{2} + … + p_{n} = 1"], size=24)
pic(s, "05_sum_one", 0.9, 4.0, 11.6, 2.6)

# 8. 공학 예시: 불량품
s = new("공학 예시: 부품 검사의 불량품 개수",
        "불량률 0.1인 공정에서 부품 3개 검사  |  X = 불량품 개수",
        notes="자체 예시 (교과서 외). 각 부품의 불량 여부가 서로 독립이라는 가정. P(X=0)=0.9³=0.729, P(X=1)=3×0.1×0.9²=0.243, P(X=2)=3×0.1²×0.9=0.027, P(X=3)=0.1³=0.001. 합계 1 확인. 다음 회차 이항분포의 예고.")
pic(s, "06_defect_pmf", 0.8, 1.9, 6.4, 4.9)
keybox(s, 7.6, 2.2, 5.1, 3.9, "읽는 법", [
    "불량품 0개일 확률이 가장 큼",
    "2개 이상일 확률 = 0.027 + 0.001",
    "  = 0.028",
    "확률의 합 = 1 확인",
])

# 9. 연속확률변수
s = new("연속확률변수: 지하철 대기 시간",
        "10분 간격 운행  |  X = 기다리는 시간 (0 ≤ X ≤ 10)",
        notes="교과서 84쪽 생각 열기. P(a≤X≤b) = (b-a)/10. 확률 = 그래프 아래 넓이라는 관점의 시작. f(x)=1/10의 그래프와 x축 사이 전체 넓이 1.")
pic(s, "07_subway_uniform", 0.8, 1.9, 7.2, 4.9)
keybox(s, 8.4, 2.2, 4.3, 3.9, "핵심", [
    "확률 = 그래프 아래 넓이",
    "P(a ≤ X ≤ b) = (b - a)/10",
    "전체 넓이 = 1",
], size=19)

# 10. 확률밀도함수의 성질
s = new("확률밀도함수의 성질", "y = f(x): 연속확률변수 X의 확률밀도함수",
        notes="교과서 85쪽 성질 1~3. 이산의 '합 = 1'이 연속에서는 '넓이 = 1'로 바뀌는 대응 관계 강조.")
pic(s, "08_pdf_area", 0.7, 1.9, 6.7, 4.9)
keybox(s, 7.8, 2.0, 4.9, 4.6, "α ≤ x ≤ β 에서", [
    "① f(x) ≥ 0",
    "② 그래프와 x축 사이 넓이 = 1",
    "③ P(a ≤ X ≤ b)",
    "    = x = a, x = b 사이의 넓이",
], size=19)

# 11. 한 점의 확률
s = new("연속확률변수에서 한 점의 확률 = 0",
        "선분의 넓이는 0  →  등호 포함 여부와 무관한 확률",
        notes="교과서 85쪽 생각과 표현. P(a≤X≤b) = P(a≤X<b) + P(X=b), P(X=b)=0. 따라서 P(a≤X≤b) = P(a<X≤b) = P(a<X<b). 이산확률변수에서는 성립하지 않는 점 주의.")
pic(s, "09_point_zero", 0.6, 1.9, 12.1, 3.2)
keybox(s, 2.2, 5.2, 8.9, 1.6, "결론", ["P(a ≤ X ≤ b) = P(a < X ≤ b) = P(a < X < b)"], size=20)

# 12. 공학 예시: 부품 수명
s = new("공학 예시: 부품 수명",
        "평균 수명 1000시간인 부품  |  X = 고장까지의 시간",
        notes="자체 예시 (교과서 외). 확률밀도함수 f(x) = (1/1000)e^(-x/1000), x ≥ 0 (지수분포, 교과서 범위 밖이므로 식은 참고용). P(X≤500) = 1 - e^(-0.5) ≈ 0.39. 신뢰성 공학의 출발점이라는 동기 부여용.")
pic(s, "10_lifetime", 0.7, 1.9, 7.4, 4.9)
keybox(s, 8.5, 2.2, 4.2, 3.9, "해석", [
    "500시간 안의 고장 확률",
    "  ≈ 0.39",
    "곡선 아래 넓이로 계산",
    "보증 기간 설정의 근거",
], size=19)

# 13. 요약
s = new("요약: 이산 vs 연속", notes="오늘 내용 정리. 합 ↔ 넓이 대응이 핵심.", source=False)
rows = [("", "이산확률변수", "연속확률변수"),
        ("취하는 값", "셀 수 있는 값", "구간 안의 모든 실수"),
        ("분포 함수", "확률질량함수 P(X = x_{i}) = p_{i}", "확률밀도함수 f(x)"),
        ("확률 계산", "확률의 합", "그래프 아래 넓이"),
        ("전체", "p_{1} + … + p_{n} = 1", "전체 넓이 = 1"),
        ("한 점의 확률", "0 이상", "항상 0"),
        ("공학 예시", "불량품 개수", "대기 시간, 부품 수명")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.9), Inches(1.7), Inches(11.5), Inches(4.9)).table
tbl.columns[0].width = Inches(2.5); tbl.columns[1].width = Inches(4.5); tbl.columns[2].width = Inches(4.5)
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if c == 0 else WHITE)
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        runs(para, v, 19, WHITE if r == 0 else (NAVY if c == 0 else DARK), r == 0 or c == 0)

# 14. 확인 퀴즈
s = new("확인 퀴즈",
        notes="정답. Q1: 0, 1, 2, 3 (교과서 81쪽 문제 1). Q2: k + 2k + 3k = 1 → k = 1/6, P(X ≥ 2) = 2k + 3k = 5/6. Q3: 넓이 (1/2)×4×4k = 1 → k = 1/8, P(0≤X≤2) = (1/2)×2×(2/8) = 1/4 (교과서 85쪽 문제 3).")
qs = [
    ("Q1", "100원짜리 동전 3개를 동시에 던질 때, 앞면의 개수 X가 취하는 모든 값"),
    ("Q2", "X의 확률분포가 P(X=1) = k, P(X=2) = 2k, P(X=3) = 3k일 때, 상수 k의 값과 P(X ≥ 2)"),
    ("Q3", "0 ≤ X ≤ 4, 확률밀도함수 f(x) = kx일 때, 상수 k의 값과 P(0 ≤ X ≤ 2)"),
]
for i, (h, q) in enumerate(qs):
    y = 1.6 + i * 1.75
    rect(s, 0.9, y, 0.95, 0.95, ORANGE, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0.9, y, 0.95, 0.95, h, size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.1, y - 0.05, 6.6 if i == 2 else 10.5, 1.3, q, size=21, anchor=MSO_ANCHOR.MIDDLE)
pic(s, "11_quiz_kx", 9.0, 4.2, 3.7, 2.7)

prs.save(OUT)
import json
json.dump(NOTES, open(OUT.rsplit(".", 1)[0] + "_notes.json", "w", encoding="utf-8"), ensure_ascii=False)
print("saved", OUT, len(prs.slides), "slides")
